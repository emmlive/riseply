"""Tailored cover letters: written from the tailored resume, never inventing facts."""
import io
import uuid
from unittest.mock import MagicMock

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app import models
from app.security import get_current_user
from app.services import cover_letter as cl, resume_customizer as rc, usage

RESUME = [
    {"type": "name", "text": "Pat Example"},
    {"type": "contact", "text": "pat@example.com  |  (312) 555-0100  |  Chicago, IL"},
    {"type": "heading", "text": "Experience"},
    {"type": "entry", "left": "Security Engineer", "right": "2019 – 2024"},
    {"type": "sub", "text": "Acme — Chicago, IL"},
    {"type": "bullet", "text": "Cut audit findings 40% across 12 business units."},
]

GOOD = ("Dear Hiring Team,\n\n" + ("I cut audit findings 40% across 12 business units at Acme. " * 6) +
        "\n\nSincerely,\nPat Example")


def _client_returning(*texts):
    c = MagicMock()
    c.messages.create.side_effect = [MagicMock(content=[MagicMock(text=t)]) for t in texts]
    return c


def test_letter_that_uses_only_resume_facts_is_accepted(monkeypatch):
    monkeypatch.setattr(rc, "client", _client_returning(GOOD))
    out = cl.generate(RESUME, {"title": "Lead", "company": "Beta", "description": "audit work"}, "Pat Example")
    assert "40%" in out and out.startswith("Dear Hiring Team")


def test_invented_number_is_retried_then_refused(monkeypatch):
    bad = GOOD.replace("40%", "87%")
    monkeypatch.setattr(rc, "client", _client_returning(bad, bad))
    with pytest.raises(ValueError):
        cl.generate(RESUME, {"title": "Lead", "company": "Beta", "description": "audit work"}, "Pat")


def test_invented_number_then_clean_retry_succeeds(monkeypatch):
    monkeypatch.setattr(rc, "client", _client_returning(GOOD.replace("40%", "87%"), GOOD))
    assert "40%" in cl.generate(RESUME, {"title": "Lead", "company": "Beta", "description": ""}, "Pat")


def test_numbers_from_the_posting_are_allowed():
    assert cl.new_numbers("the 5 teams you lead", "none", "lead 5 teams") == []


def test_docx_has_letterhead_and_paragraphs():
    data = cl.build_docx(GOOD, "Pat Example", "pat@example.com")
    paras = [p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]
    assert paras[0] == "Pat Example" and paras[1] == "pat@example.com"
    assert any(p.startswith("Dear Hiring Team") for p in paras)
    assert paras[-1].startswith("Sincerely")


@pytest.fixture()
def setup():
    db = SessionLocal()
    user = models.User(email=f"cl-{uuid.uuid4().hex[:8]}@example.com", hashed_password="x", resume_text="r")
    db.add(user)
    db.commit()
    db.refresh(user)
    job = models.Job(source="greenhouse", external_id=uuid.uuid4().hex, company="Beta Corp", title="Lead",
                     description="audit work", url="https://x.example/2")
    db.add(job)
    db.commit()
    application = models.Application(user_id=user.id, job_id=job.id, status="pending_approval",
                                      tailored_resume_path="Beta.docx",
                                      tailored_resume_data=rc.build_docx_from_blocks(RESUME))
    db.add(application)
    db.commit()
    db.refresh(application)
    app.dependency_overrides[get_current_user] = lambda: user
    yield db, user, application, TestClient(app)
    app.dependency_overrides.clear()
    db.close()


def test_write_save_download_flow_and_metering(setup, monkeypatch):
    db, user, application, client = setup
    monkeypatch.setattr(rc, "client", _client_returning(GOOD))
    assert client.get(f"/applications/{application.id}/cover-letter").json()["text"] == ""
    r = client.post(f"/applications/{application.id}/cover-letter")
    assert r.status_code == 200 and "40%" in r.json()["text"]
    assert usage.get_usage(db, user.id, "interview_prep") == 1

    r = client.put(f"/applications/{application.id}/cover-letter", json={"text": GOOD + "\nP.S. Thanks."})
    assert r.status_code == 200
    assert client.put(f"/applications/{application.id}/cover-letter", json={"text": "  "}).status_code == 422

    d = client.get(f"/applications/{application.id}/cover-letter/download")
    assert d.status_code == 200 and "Cover_Letter_Beta_Corp.docx" in d.headers["content-disposition"]
    assert Document(io.BytesIO(d.content)).paragraphs[0].text == "Pat Example"

    lst = client.get("/applications").json()
    assert any(a["id"] == application.id and a["has_cover_letter"] for a in lst)


def test_failed_writing_is_not_charged(setup, monkeypatch):
    db, user, application, client = setup
    bad = GOOD.replace("40%", "87%")
    monkeypatch.setattr(rc, "client", _client_returning(bad, bad))
    r = client.post(f"/applications/{application.id}/cover-letter")
    assert r.status_code == 502
    assert usage.get_usage(db, user.id, "interview_prep") == 0


def test_other_users_letter_is_not_reachable(setup):
    db, user, application, client = setup
    app.dependency_overrides[get_current_user] = lambda: models.User(id=999999, email="x@y.z", hashed_password="x")
    assert client.get(f"/applications/{application.id}/cover-letter").status_code == 404
