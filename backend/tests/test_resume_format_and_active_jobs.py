"""Tailored resume layout + preview, and the 'only open postings' rules."""
import json
import uuid
from datetime import datetime, timedelta
from unittest.mock import patch, MagicMock

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app import models
from app.security import get_current_user
from app.services import pipeline_runner, posting_check, resume_customizer as rc

_n = [0]

SAMPLE = {
    "name": "Jane Doe", "headline": "Security Auditor",
    "contact": ["jane@x.com", "Plainfield, IL"],
    "sections": [
        {"heading": "Professional Summary", "kind": "text", "text": "Auditor."},
        {"heading": "Core Skills", "kind": "skills", "skills": [{"label": "Frameworks", "items": ["SOC 2", "NIST"]}]},
        {"heading": "Experience", "kind": "entries", "entries": [
            {"title": "IT Auditor", "org": "Fed", "location": "Chicago, IL", "dates": "2019 – Present", "bullets": ["Led audits.", "Cut findings 30%."]}]},
        {"heading": "Certifications", "kind": "bullets", "bullets": ["CISA"]},
    ],
}


@pytest.fixture()
def db():
    s = SessionLocal()
    yield s
    s.close()


@pytest.fixture(autouse=True)
def _clear():
    yield
    app.dependency_overrides.clear()


def _user(db):
    _n[0] += 1
    u = models.User(email=f"fmt{_n[0]}@x.com", hashed_password="x", resume_text="Jane Doe\nAuditor.")
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _client_as(user_id):
    from app.database import get_db
    from fastapi import Depends

    def _u(db=Depends(get_db)):
        return db.get(models.User, user_id)
    app.dependency_overrides[get_current_user] = _u
    return TestClient(app)


# --- resume document ------------------------------------------------------

def test_structure_becomes_blocks_in_order():
    blocks = rc.structure_to_blocks(SAMPLE)
    assert [b["type"] for b in blocks] == [
        "name", "headline", "contact",
        "heading", "text", "heading", "skills",
        "heading", "entry", "sub", "bullet", "bullet", "heading", "bullet",
    ]
    assert blocks[2]["text"] == "jane@x.com  |  Plainfield, IL"
    assert blocks[8] == {"type": "entry", "left": "IT Auditor", "right": "2019 – Present"}
    assert blocks[9]["text"] == "Fed — Chicago, IL"


def test_docx_round_trips_to_the_same_blocks():
    blocks = rc.structure_to_blocks(SAMPLE)
    again = rc.docx_to_blocks(rc.build_docx_from_blocks(blocks))
    assert again == blocks


def test_old_style_docx_still_previews():
    from docx import Document
    import io
    d = Document()
    d.add_heading("Jane Doe", level=2)
    d.add_heading("Experience", level=2)
    d.add_paragraph("Did things", style="List Bullet")
    d.add_paragraph("Plain line")
    buf = io.BytesIO()
    d.save(buf)
    types = [b["type"] for b in rc.docx_to_blocks(buf.getvalue())]
    assert types == ["name", "heading", "bullet", "text"]


def test_tailor_uses_json_and_falls_back_to_text():
    def reply(text):
        m = MagicMock()
        m.content = [MagicMock(text=text)]
        return m

    job = {"title": "T", "company": "C", "description": "d"}
    good = json.dumps({"rationale": "Moved audit work up.", "resume": SAMPLE})
    with patch.object(rc.client.messages, "create", return_value=reply(good)):
        rationale, blocks = rc.tailor_resume("base", job)
    assert rationale == "Moved audit work up."
    assert blocks[0] == {"type": "name", "text": "Jane Doe"}

    messy = "RATIONALE:\nKept it.\n---RESUME---\nJANE DOE\nEXPERIENCE\n- Led audits"
    with patch.object(rc.client.messages, "create", return_value=reply(messy)):
        rationale, blocks = rc.tailor_resume("base", job)
    assert rationale == "Kept it."
    assert [b["type"] for b in blocks] == ["name", "heading", "bullet"]


def test_preview_endpoint_returns_blocks_and_is_private(db):
    user, other = _user(db), _user(db)
    job = models.Job(source="t", external_id=uuid.uuid4().hex, company="Acme", title="Auditor",
                     description="d", discovered_at=datetime.utcnow())
    db.add(job)
    db.commit()
    a = models.Application(user_id=user.id, job_id=job.id, status="pending_approval",
                           tailored_resume_path="Acme.docx", tailoring_rationale="why",
                           tailored_resume_data=rc.build_docx_from_blocks(rc.structure_to_blocks(SAMPLE)))
    empty = models.Application(user_id=user.id, job_id=job.id, status="pending_approval")
    db.add_all([a, empty])
    db.commit()

    client = _client_as(user.id)
    body = client.get(f"/applications/{a.id}/tailored-resume/preview").json()
    assert body["filename"] == "Acme.docx" and body["rationale"] == "why"
    assert body["blocks"][0]["text"] == "Jane Doe"
    assert client.get(f"/applications/{empty.id}/tailored-resume/preview").status_code == 404

    client = _client_as(other.id)
    assert client.get(f"/applications/{a.id}/tailored-resume/preview").status_code == 404


# --- is the posting still open? -------------------------------------------

def _fake_response(status=200, body=b"", headers=None):
    r = MagicMock()
    r.status_code = status
    r.headers = headers or {}
    r.encoding = "utf-8"
    r.iter_content.return_value = [body]
    return r


def test_check_posting_verdicts():
    with patch.object(posting_check, "_is_public_host", return_value=True):
        with patch("requests.get", return_value=_fake_response(404)):
            assert posting_check.check_posting("https://x.com/j/1") is False
        with patch("requests.get", return_value=_fake_response(200, b"<p>Sorry, this job is no longer available.</p>")):
            assert posting_check.check_posting("https://x.com/j/1") is False
        with patch("requests.get", return_value=_fake_response(200, b"<h1>Security Engineer</h1> Apply now")):
            assert posting_check.check_posting("https://x.com/j/1") is True
        with patch("requests.get", return_value=_fake_response(403)):
            assert posting_check.check_posting("https://x.com/j/1") is None  # blocked: unknown, keep
        with patch("requests.get", side_effect=OSError("boom")):
            assert posting_check.check_posting("https://x.com/j/1") is None
        redirect = _fake_response(302, headers={"Location": "https://boards.greenhouse.io/acme?error=true"})
        with patch("requests.get", return_value=redirect):
            assert posting_check.check_posting("https://x.com/j/1") is False


def test_check_posting_refuses_private_addresses():
    assert posting_check.check_posting("http://127.0.0.1/admin") is None
    assert posting_check.check_posting("http://169.254.169.254/latest") is None
    assert posting_check.check_posting("") is None
    assert posting_check.check_posting("file:///etc/passwd") is None


def _profile_user(db):
    u = _user(db)
    db.add(models.SearchProfile(
        user_id=u.id, name="P", titles=json.dumps(["Security"]), locations=json.dumps([]),
        seniority=json.dumps([]), min_match_score=60, exclude_companies=json.dumps([]),
        keywords_required=json.dumps([]), keywords_excluded=json.dumps([]), active=True))
    db.commit()
    return u


def _job(db, company, **kw):
    j = models.Job(source="test", external_id=uuid.uuid4().hex, company=company, title=f"Security {company}",
                   location="Remote", description="d", url=f"https://{company}.example/job",
                   discovered_at=datetime.utcnow(), **kw)
    db.add(j)
    db.commit()
    return j


def test_matching_skips_closed_stale_and_dead_links(db):
    db.query(models.Job).delete()
    db.query(models.ScoredJob).delete()
    db.commit()
    user = _profile_user(db)
    live = _job(db, "Live")
    closed_flag = _job(db, "Flagged", is_active=False)
    old = _job(db, "Old", posted_at=datetime.utcnow() - timedelta(days=90))
    dead_link = _job(db, "Dead")
    unknown = _job(db, "Blocked")

    def score(job, resume, profiles, ignore_location=False):
        return {"profile_name": "P", "score": 90, "reason": "r", "meets_threshold": True}

    def verdict(url):
        return False if "Dead" in url else (None if "Blocked" in url else True)

    scored = []
    with patch("app.services.pipeline_runner.matcher.best_profile_match", side_effect=lambda *a, **k: (scored.append(a[0]["company"]), score(*a, **k))[1]), \
         patch("app.services.pipeline_runner.posting_check.check_posting", side_effect=verdict), \
         patch("app.services.pipeline_runner.resume_customizer.customize_for_job", return_value=("r.docx", b"x", "why")), \
         patch("app.services.pipeline_runner.notifier.notify_new_match"):
        result = pipeline_runner.run_matching_for_user(db, user, skip_usage_metering=True)

    assert sorted(scored) == ["Blocked", "Dead", "Live"]  # flagged-closed and 90-day-old never scored
    queued_jobs = {db.get(models.Application, i).job_id for i in result["queued_application_ids"]}
    assert queued_jobs == {live.id, unknown.id}            # the dead link was dropped, the unknown one kept
    db.refresh(dead_link)
    assert dead_link.is_active is False                     # and won't be offered to anyone else


def test_discovery_closes_jobs_missing_from_their_company_board(db):
    db.query(models.Job).delete()
    db.commit()
    gone = models.Job(source="greenhouse", external_id="1", company="acme", title="Old role",
                      discovered_at=datetime.utcnow() - timedelta(days=5), last_seen_at=datetime.utcnow() - timedelta(days=5))
    still = models.Job(source="greenhouse", external_id="2", company="acme", title="Open role",
                       discovered_at=datetime.utcnow() - timedelta(days=5), last_seen_at=datetime.utcnow() - timedelta(days=5))
    other_company = models.Job(source="greenhouse", external_id="3", company="failedco", title="Unknown",
                               discovered_at=datetime.utcnow() - timedelta(days=5))
    db.add_all([gone, still, other_company])
    db.commit()

    listing = [{"source": "greenhouse", "external_id": "2", "company": "acme", "title": "Open role",
                "location": "", "url": "", "description": ""}]
    empty = []
    with patch("app.services.pipeline_runner.greenhouse.fetch_all", return_value=listing), \
         patch("app.services.pipeline_runner.lever.fetch_all", return_value=empty), \
         patch("app.services.pipeline_runner.rss_boards.fetch_all", return_value=[]), \
         patch("app.services.pipeline_runner.remoteok.fetch_jobs", return_value=[]), \
         patch("app.services.pipeline_runner.arbeitnow.fetch_jobs", return_value=[]), \
         patch("app.services.pipeline_runner.adzuna.fetch_by_keywords", return_value=[]), \
         patch("app.services.pipeline_runner.adzuna.fetch_by_keyword_location_pairs", return_value=[], create=True), \
         patch("app.services.pipeline_runner.usajobs.fetch_by_keywords", return_value=[]):
        pipeline_runner.run_discovery(db)

    for j in (gone, still, other_company):
        db.refresh(j)
    assert gone.is_active is False         # no longer on acme's board
    assert still.is_active is True and still.last_seen_at > datetime.utcnow() - timedelta(minutes=5)
    assert other_company.is_active is not False  # company wasn't fetched, so nothing is assumed


# --- editing a tailored resume -------------------------------------------

def _app_with_resume(db, user):
    job = models.Job(source="t", external_id=uuid.uuid4().hex, company="Acme", title="Auditor",
                     description="d", discovered_at=datetime.utcnow())
    db.add(job)
    db.commit()
    a = models.Application(user_id=user.id, job_id=job.id, status="pending_approval", tailored_resume_path="Acme.docx",
                           tailored_resume_data=rc.build_docx_from_blocks(rc.structure_to_blocks(SAMPLE)))
    db.add(a)
    db.commit()
    return a


def test_saving_edits_rebuilds_the_document(db):
    user = _user(db)
    a = _app_with_resume(db, user)
    client = _client_as(user.id)

    blocks = client.get(f"/applications/{a.id}/tailored-resume/preview").json()["blocks"]
    edited = [b for b in blocks if b.get("text") != "Cut findings 30%."]       # remove a bullet
    edited.append({"type": "heading", "text": "Projects"})                      # add a section
    edited.append({"type": "bullet", "text": "Built a SOC 2 readiness tracker."})
    edited.append({"type": "bullet", "text": "   "})                            # blank rows are dropped

    res = client.put(f"/applications/{a.id}/tailored-resume", json={"blocks": edited})
    assert res.status_code == 200
    texts = [b.get("text") or b.get("left") for b in res.json()["blocks"]]
    assert "Cut findings 30%." not in texts
    assert texts[-2:] == ["Projects", "Built a SOC 2 readiness tracker."]

    db.expire_all()
    stored = rc.docx_to_blocks(db.get(models.Application, a.id).tailored_resume_data)
    assert stored == res.json()["blocks"]                                      # the download matches what was saved


def test_saving_edits_validation_and_ownership(db):
    user, other = _user(db), _user(db)
    a = _app_with_resume(db, user)
    client = _client_as(user.id)
    url = f"/applications/{a.id}/tailored-resume"

    assert client.put(url, json={"blocks": []}).status_code == 422
    assert client.put(url, json={"blocks": [{"type": "bullet", "text": " "}]}).status_code == 422
    assert client.put(url, json={"blocks": [{"type": "evil", "text": "x"}]}).status_code == 422
    assert client.put(url, json={"blocks": [{"type": "text", "text": "x"}] * 401}).status_code == 422

    client = _client_as(other.id)
    assert client.put(url, json={"blocks": [{"type": "text", "text": "hijack"}]}).status_code == 404


def test_legacy_resume_gets_a_name_and_contact_line():
    from docx import Document
    import io
    d = Document()
    d.add_heading("Jane Doe", level=2)
    d.add_paragraph("Security Architect | GRC")
    d.add_paragraph("Plainfield, IL • (555) 123-4567 • jane@x.com")
    d.add_heading("Summary", level=2)
    d.add_paragraph("Auditor.")
    buf = io.BytesIO()
    d.save(buf)
    types = [b["type"] for b in rc.docx_to_blocks(buf.getvalue())]
    assert types == ["name", "headline", "contact", "heading", "text"]
