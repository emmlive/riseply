"""Task guide: building, saving progress, help, ownership, metering and parsing.
The model call is patched; these cover the plumbing and the validation of what
the model returns."""
import json
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app.security import get_current_user
from app import models
from app.services import usage, task_guide

client = TestClient(app)
_n = [0]
SVC = "app.routers.task_guides.task_guide_service"

GOOD = {
    "title": "Weekly status report", "goal": "A one-page report your manager can read in two minutes.",
    "deliverable": "a one-page report", "why_it_matters": "It drives Monday's meeting.",
    "confirm_first": ["Who reads it?", "What format?"],
    "requirements": [{"item": "Last week's numbers", "why": "To compare"}, "Access to the tracker"],
    "materials": [{"name": "Report template", "purpose": "Skeleton", "template": "# Status\n- [DONE]\n- [NEXT]"}],
    "steps": [
        {"title": "Collect numbers", "what_to_do": "Pull the figures", "how": "From the tracker", "done_when": "You have 5 numbers", "minutes": 30, "watch_out": "Old data"},
        {"title": "Write it up", "what_to_do": "Fill the template", "how": "Short sentences", "done_when": "One page", "minutes": "45", "watch_out": ""},
    ],
    "quality_check": ["Fits one page?"], "common_mistakes": ["Too long"], "total_time": "about 90 minutes",
}


def _user(db):
    _n[0] += 1
    u = models.User(email=f"tg{_n[0]}@x.com", hashed_password="x", full_name="T", resume_text="Analyst")
    db.add(u); db.commit(); db.refresh(u)
    return u


@pytest.fixture()
def db():
    s = SessionLocal()
    yield s
    s.close()


@pytest.fixture(autouse=True)
def _clear():
    yield
    app.dependency_overrides.clear()


def _as(u):
    app.dependency_overrides[get_current_user] = lambda: u


def _build(u, built=None):
    _as(u)
    with patch(f"{SVC}.build", return_value=built or task_guide.parse(json.dumps(GOOD))):
        return client.post("/task-guides", json={"task": "Write the weekly status report for my team", "role": "Analyst", "level": "first_time"})


def test_parse_normalises_and_numbers_steps():
    d = task_guide.parse("```json\n" + json.dumps(GOOD) + "\n```")
    assert [s["id"] for s in d["steps"]] == ["s1", "s2"]
    assert d["steps"][1]["minutes"] == 45
    assert d["requirements"][1] == {"item": "Access to the tracker", "why": ""}
    assert d["materials"][0]["template"].startswith("# Status")


def test_parse_rejects_thin_guides():
    with pytest.raises(ValueError):
        task_guide.parse(json.dumps({"goal": "x" * 10, "steps": [{"title": "only one"}]}))
    with pytest.raises(ValueError):
        task_guide.parse("not json")


def test_build_saves_and_counts(db):
    u = _user(db)
    r = _build(u)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["title"] == "Weekly status report" and body["level"] == "first_time"
    assert len(body["steps"]) == 2 and body["done"] == []
    db.expire_all()
    assert usage.get_usage(db, u.id, "task_guide") == 1
    assert client.get("/task-guides").json()[0]["steps"] == 2


def test_failed_build_is_not_counted(db):
    u = _user(db)
    _as(u)
    with patch(f"{SVC}.build", side_effect=RuntimeError("boom")):
        r = client.post("/task-guides", json={"task": "Write the weekly status report for my team"})
    assert r.status_code == 502
    db.expire_all()
    assert usage.get_usage(db, u.id, "task_guide") == 0


def test_progress_keeps_only_real_steps(db):
    u = _user(db)
    gid = _build(u).json()["id"]
    r = client.put(f"/task-guides/{gid}/progress", json={"done": ["s2", "s2", "s9", "x"]})
    assert r.json()["done"] == ["s2"]
    assert client.get("/task-guides").json()[0]["done"] == 1


def test_help_and_metering(db):
    u = _user(db)
    gid = _build(u).json()["id"]
    with patch(f"{SVC}.help_reply", return_value="Looks good; add the dates.") as h:
        r = client.post(f"/task-guides/{gid}/help", json={"step_id": "s1", "message": "Here is my draft: ...", "history": []})
    assert r.json() == {"reply": "Looks good; add the dates."}
    assert h.called
    db.expire_all()
    assert usage.get_usage(db, u.id, "coaching_message") == 1
    with patch(f"{SVC}.help_reply", side_effect=RuntimeError("x")):
        assert client.post(f"/task-guides/{gid}/help", json={"step_id": "s1", "message": "hello there"}).status_code == 502
    db.expire_all()
    assert usage.get_usage(db, u.id, "coaching_message") == 1


def test_other_users_cannot_see_or_change_a_guide(db):
    a, b = _user(db), _user(db)
    gid = _build(a).json()["id"]
    _as(b)
    assert client.get(f"/task-guides/{gid}").status_code == 404
    assert client.put(f"/task-guides/{gid}/progress", json={"done": []}).status_code == 404
    assert client.delete(f"/task-guides/{gid}").status_code == 404
    assert client.get("/task-guides").json() == []


def test_delete(db):
    u = _user(db)
    gid = _build(u).json()["id"]
    assert client.delete(f"/task-guides/{gid}").json() == {"ok": True}
    assert client.get(f"/task-guides/{gid}").status_code == 404


def test_task_text_must_be_real():
    r = client.post("/task-guides", json={"task": "hi"})
    assert r.status_code in (401, 403, 422)


def test_help_reply_builds_a_clean_conversation():
    fake = MagicMock()
    fake.messages.create.return_value = MagicMock(content=[MagicMock(text="ok")])
    guide = task_guide.parse(json.dumps(GOOD))
    with patch.object(task_guide, "client", fake):
        out = task_guide.help_reply("task", guide, "s1", "my draft", [{"role": "assistant", "content": "hi"}, {"role": "user", "content": "q"}], "first_time")
    assert out == "ok"
    kw = fake.messages.create.call_args.kwargs
    assert kw["messages"][0]["role"] == "user" and kw["messages"][-1]["content"] == "my draft"
    assert "Collect numbers" in kw["system"]


# ---- exports ----
from app.services import doc_export

MD = """# Requirements Outline
## Controls audit
**Version:** [1.0]
**Date:** [MM/DD/YYYY]
---
## 1. Purpose
Why this exists, with **bold** words.
## 2. Roles
| Role | Responsibility |
|---|---|
| Lead | Runs the audit |
| Owner | Uploads **evidence** |
## 3. To do
- [ ] First
- [x] Second
1. One
2. Two
"""


def test_parse_markdown_blocks():
    kinds = [b["type"] for b in doc_export.parse(MD)]
    assert kinds == ["heading", "heading", "para", "para", "rule", "heading", "para", "heading", "table", "heading", "bullets", "numbers"]
    table = [b for b in doc_export.parse(MD) if b["type"] == "table"][0]
    assert table["header"] == ["Role", "Responsibility"] and table["rows"][1] == ["Owner", "Uploads **evidence**"]
    checks = [b for b in doc_export.parse(MD) if b["type"] == "bullets"][0]["items"]
    assert [c["check"] for c in checks] == [False, True]
    assert doc_export.has_table(MD) and not doc_export.has_table("# Only text")


def test_docx_pdf_xlsx_are_real_files():
    import io
    from docx import Document
    from openpyxl import load_workbook
    blocks = doc_export.parse(MD)
    d = Document(io.BytesIO(doc_export.to_docx(blocks, "T")))
    assert len(d.tables) == 1 and d.tables[0].rows[1].cells[0].text == "Lead"
    assert any(p.text == "Requirements Outline" for p in d.paragraphs)
    assert doc_export.to_pdf(blocks, "T").startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(doc_export.to_xlsx(blocks, "T")))
    assert wb.sheetnames[0].startswith("2. Roles") and wb.worksheets[0]["A1"].value == "Role"
    assert wb.worksheets[0]["B3"].value == "Uploads evidence"
    assert "Notes" in wb.sheetnames


def test_xlsx_without_a_table_is_one_sheet():
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(doc_export.to_xlsx(doc_export.parse("# Email\nHello\n- a\n- b"), "Email")))
    assert len(wb.sheetnames) == 1


def test_guide_markdown_and_share_text():
    g = task_guide.parse(json.dumps(GOOD))
    md = doc_export.guide_markdown(g, ["s1"], include_materials=True)
    assert "- [x] Pull the figures" in md and "- [ ] Fill the template" in md
    assert "## Template: Report template" in md and "### Status" in md
    t = doc_export.share_text(g, ["s2"])
    assert "[ ] 1. Collect numbers" in t and "[x] 2. Write it up" in t and "Done when: One page" in t


def test_export_endpoints(db):
    u = _user(db)
    built = task_guide.parse(json.dumps({**GOOD, "materials": [{"name": "Roles table", "kind": "spreadsheet", "purpose": "p", "template": "# Roles\n| A | B |\n|---|---|\n| 1 | 2 |"}]}))
    gid = _build(u, built).json()["id"]
    client.put(f"/task-guides/{gid}/progress", json={"done": ["s1"]})
    r = client.get(f"/task-guides/{gid}/export?format=docx")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "attachment" in r.headers["content-disposition"]
    assert client.get(f"/task-guides/{gid}/export?format=pdf").content.startswith(b"%PDF")
    md = client.get(f"/task-guides/{gid}/export?format=md&include_templates=true").text
    assert "- [x] Pull the figures" in md and "Template: Roles table" in md
    x = client.get(f"/task-guides/{gid}/materials/0/export?format=xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert client.get(f"/task-guides/{gid}/materials/5/export?format=docx").status_code == 404
    assert client.get(f"/task-guides/{gid}/export?format=exe").status_code == 422
    assert "[x] 1. Collect numbers" in client.get(f"/task-guides/{gid}/share-text").json()["text"]


def test_exports_are_private(db):
    a, b = _user(db), _user(db)
    gid = _build(a).json()["id"]
    _as(b)
    assert client.get(f"/task-guides/{gid}/export?format=docx").status_code == 404
    assert client.get(f"/task-guides/{gid}/materials/0/export?format=docx").status_code == 404
    assert client.get(f"/task-guides/{gid}/share-text").status_code == 404


def test_material_kind_is_validated():
    d = task_guide.parse(json.dumps({**GOOD, "materials": [{"name": "A", "kind": "weird", "template": "# x\ntext here"}, {"name": "B", "kind": "spreadsheet", "template": "| a |\n|---|\n| 1 |"}]}))
    assert [m["kind"] for m in d["materials"]] == ["document", "spreadsheet"]
