"""Text out of a file someone attaches to their own work: a draft document,
spreadsheet, PDF or plain text. Used to review a task draft against its guide.
Limits keep memory small on a small server: size, pages, sheets and rows."""
import io

MAX_BYTES = 3 * 1024 * 1024
MAX_CHARS = 24000
MAX_PDF_PAGES = 15
MAX_SHEETS = 5
MAX_ROWS = 200


class DraftError(ValueError):
    """A problem the person can fix, worded for them."""


def _docx(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise DraftError("Couldn't read that .docx file. It may be damaged. Try saving it again, or paste the text instead.")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables[:10]:
        for row in t.rows[:60]:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise DraftError("Couldn't read that spreadsheet. It may be damaged or password-protected. Try saving it again, or paste the text instead.")
    out = []
    for ws in wb.worksheets[:MAX_SHEETS]:
        out.append(f"[Sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                out.append("(more rows not shown)")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append("\t".join(cells).rstrip())
    wb.close()
    return "\n".join(out)


def _pdf(data: bytes) -> str:
    import pdfplumber
    try:
        pages = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:MAX_PDF_PAGES]:
                pages.append(page.extract_text() or "")
                page.flush_cache()
        return "\n".join(pages)
    except Exception:
        raise DraftError("Couldn't read that PDF. It may be scanned, damaged or password-protected. Paste the text instead.")


def extract(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise DraftError("That file is too large. Please keep it under 3 MB, or paste the important part.")
    name = (filename or "").lower()
    if name.endswith(".docx"):
        text = _docx(data)
    elif name.endswith((".xlsx", ".xlsm")):
        text = _xlsx(data)
    elif name.endswith(".pdf"):
        text = _pdf(data)
    elif name.endswith((".txt", ".md", ".csv")):
        text = data.decode("utf-8", errors="replace")
    elif name.endswith((".doc", ".xls")):
        raise DraftError("Old .doc and .xls files can't be read. Save it as .docx or .xlsx, or paste the text.")
    else:
        raise DraftError("Attach a .docx, .xlsx, .pdf, .txt, .md or .csv file, or paste the text.")
    text = text.replace("\x00", "").strip()
    if not text:
        raise DraftError("No text found in that file. If it is a scan or an image, paste the text instead.")
    return text[:MAX_CHARS]
