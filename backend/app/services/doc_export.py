"""Turn the plain markdown a guide or template is written in into real files:
Word (.docx), PDF, Excel (.xlsx) and the markdown itself.

Only the small subset of markdown the guide produces is understood: #/##/###
headings, paragraphs, - and 1. lists (with [ ] / [x] checkboxes), | tables |,
---, and **bold**. Anything else is kept as plain text.
"""
import io
import re

ACCENT = "1F7A6C"
INK = "16233D"

_ROW = re.compile(r"^\s*\|.*\|\s*$")
_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse(md: str) -> list[dict]:
    """Blocks: heading(level,text), para(text), bullets(items), numbers(items),
    table(header, rows), rule. List items are {text, check: None|bool}."""
    lines = (md or "").replace("\r\n", "\n").split("\n")
    out: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            out.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            out.append({"type": "rule"})
            i += 1
            continue
        if _ROW.match(s) and i + 1 < len(lines) and _SEP.match(lines[i + 1].strip()) and "|" in lines[i + 1]:
            header = _cells(s)
            rows = []
            i += 2
            while i < len(lines) and _ROW.match(lines[i].strip()):
                r = _cells(lines[i])
                rows.append((r + [""] * len(header))[:len(header)])
                i += 1
            out.append({"type": "table", "header": header, "rows": rows})
            continue
        m = re.match(r"^[-*•]\s+(.*)$", s)
        n = re.match(r"^\d+[.)]\s+(.*)$", s)
        if m or n:
            kind = "bullets" if m else "numbers"
            items = []
            while i < len(lines):
                t = lines[i].strip()
                mm = re.match(r"^[-*•]\s+(.*)$", t) if kind == "bullets" else re.match(r"^\d+[.)]\s+(.*)$", t)
                if not mm:
                    break
                text, check = mm.group(1), None
                c = re.match(r"^\[( |x|X)\]\s+(.*)$", text)
                if c:
                    check, text = c.group(1).lower() == "x", c.group(2)
                items.append({"text": text.strip(), "check": check})
                i += 1
            out.append({"type": kind, "items": items})
            continue
        out.append({"type": "para", "text": s})
        i += 1
    return out


def has_table(md: str) -> bool:
    return any(b["type"] == "table" for b in parse(md))


def _runs(text: str) -> list[tuple[str, bool]]:
    """Split on **bold** into (text, bold) runs; also drops `code` ticks."""
    text = text.replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") and len(p) > 4 else (p, False) for p in parts if p]


def _plain(text: str) -> str:
    return "".join(t for t, _ in _runs(text))


def safe_name(name: str, ext: str) -> str:
    base = re.sub(r"[^\w\-]+", "_", name or "document").strip("_")[:60] or "document"
    return f"{base}.{ext}"


# ---------- Word ----------

def to_docx(blocks: list[dict], title: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for lvl, size in ((1, 20), (2, 14), (3, 12), (4, 11)):
        st = doc.styles[f"Heading {min(lvl, 4)}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(ACCENT if lvl > 1 else INK)
    if title:
        doc.core_properties.title = title[:200]

    def add_runs(par, text, size=None, bold_all=False, color=None):
        for t, b in _runs(text):
            r = par.add_run(t)
            r.bold = b or bold_all
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    for b in blocks:
        t = b["type"]
        if t == "heading":
            h = doc.add_heading(level=min(b["level"], 4))
            add_runs(h, b["text"])
        elif t == "para":
            add_runs(doc.add_paragraph(), b["text"])
        elif t in ("bullets", "numbers"):
            for it in b["items"]:
                p = doc.add_paragraph(style="List Bullet" if t == "bullets" else "List Number")
                p.paragraph_format.space_after = Pt(2)
                if it["check"] is not None:
                    p.add_run("☑ " if it["check"] else "☐ ")
                add_runs(p, it["text"])
        elif t == "rule":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "C7CCD6")):
                bottom.set(qn(k), v)
            bdr.append(bottom)
            pPr.append(bdr)
        elif t == "table":
            cols = max(1, len(b["header"]))
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Table Grid"
            for j, h in enumerate(b["header"]):
                cell = table.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], h, size=10, bold_all=True, color="FFFFFF")
                shade(cell, ACCENT)
            for row in b["rows"]:
                cells = table.add_row().cells
                for j in range(cols):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "", size=10)
            doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Excel ----------

def to_xlsx(blocks: list[dict], title: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="C7CCD6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    used: set[str] = set()

    def sheet_name(raw: str, fallback: str) -> str:
        name = re.sub(r"[\[\]\*\?/\\:]", " ", _plain(raw)).strip()[:28] or fallback
        base, n = name, 2
        while name.lower() in used:
            name = f"{base[:25]} {n}"
            n += 1
        used.add(name.lower())
        return name

    heading = ""
    tables = 0
    notes: list[str] = []
    for b in blocks:
        if b["type"] == "heading":
            heading = b["text"]
            notes.append(_plain(b["text"]))
        elif b["type"] == "table":
            tables += 1
            ws = wb.create_sheet(sheet_name(heading or title, f"Table {tables}"))
            ws.append([_plain(h) for h in b["header"]])
            for r in b["rows"]:
                ws.append([_plain(c) for c in r])
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=ACCENT)
            for row in ws.iter_rows():
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    c.border = border
            for j in range(1, len(b["header"]) + 1):
                longest = max([len(str(ws.cell(row=r, column=j).value or "")) for r in range(1, ws.max_row + 1)] or [8])
                ws.column_dimensions[get_column_letter(j)].width = max(12, min(50, longest + 2))
            ws.freeze_panes = "A2"
            if b["rows"]:
                ws.auto_filter.ref = ws.dimensions
        elif b["type"] == "para":
            notes.append(_plain(b["text"]))
        elif b["type"] in ("bullets", "numbers"):
            notes += [("☐ " if i["check"] is False else "☑ " if i["check"] else "• ") + _plain(i["text"]) for i in b["items"]]
    if tables == 0 or notes and tables:
        ws = wb.create_sheet(sheet_name("Notes" if tables else (title or "Template"), "Template"))
        for line in notes:
            ws.append([line])
        ws.column_dimensions["A"].width = 100
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

def to_pdf(blocks: list[dict], title: str = "") -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink, accent, rule = HexColor("#" + INK), HexColor("#" + ACCENT), HexColor("#C7CCD6")

    def st(name, size, **kw):
        return ParagraphStyle(name, fontName=kw.pop("font", "Helvetica"), fontSize=size, leading=kw.pop("leading", size * 1.3), textColor=kw.pop("color", ink), **kw)

    S = {
        1: st("h1", 20, font="Helvetica-Bold", spaceAfter=6, spaceBefore=2),
        2: st("h2", 14, font="Helvetica-Bold", color=accent, spaceBefore=12, spaceAfter=4, keepWithNext=1),
        3: st("h3", 11.5, font="Helvetica-Bold", color=accent, spaceBefore=8, spaceAfter=3, keepWithNext=1),
        4: st("h4", 10.5, font="Helvetica-Bold", spaceBefore=6, spaceAfter=2, keepWithNext=1),
        "p": st("p", 10, spaceAfter=5, leading=13.5),
        "li": st("li", 10, spaceAfter=2, leftIndent=16, bulletIndent=3, leading=13.5),
        "cell": st("cell", 9, leading=11.5),
        "head": st("head", 9, font="Helvetica-Bold", color=white, leading=11.5),
    }

    def rich(text: str) -> str:
        return "".join(f"<b>{escape(t)}</b>" if b else escape(t) for t, b in _runs(text))

    story = []
    for b in blocks:
        t = b["type"]
        if t == "heading":
            story.append(Paragraph(rich(b["text"]), S[min(b["level"], 4)]))
        elif t == "para":
            story.append(Paragraph(rich(b["text"]), S["p"]))
        elif t in ("bullets", "numbers"):
            for n, it in enumerate(b["items"], 1):
                mark = ("☑" if it["check"] else "☐") if it["check"] is not None else ("•" if t == "bullets" else f"{n}.")
                story.append(Paragraph(rich(it["text"]), S["li"], bulletText=mark if it["check"] is None else "[x]" if it["check"] else "[ ]"))
        elif t == "rule":
            story.append(HRFlowable(width="100%", thickness=0.6, color=rule, spaceBefore=4, spaceAfter=6))
        elif t == "table":
            cols = max(1, len(b["header"]))
            data = [[Paragraph(rich(h), S["head"]) for h in b["header"]]]
            data += [[Paragraph(rich(c), S["cell"]) for c in (r + [""] * cols)[:cols]] for r in b["rows"]]
            tbl = Table(data, colWidths=[(letter[0] - 1.8 * inch) / cols] * cols, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), accent),
                ("GRID", (0, 0), (-1, -1), 0.4, rule),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white, HexColor("#F5F6F8")]),
            ]))
            story += [tbl, Spacer(1, 8)]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch, title=(title or "Document")[:200])
    doc.build(story or [Spacer(1, 1)])
    return buf.getvalue()


# ---------- the whole guide as one document ----------

def guide_markdown(g: dict, done: list[str] | None = None, include_materials: bool = False) -> str:
    done = set(done or [])
    L = [f"# {g.get('title') or 'Task guide'}", ""]
    if g.get("goal"):
        L += [f"**Goal:** {g['goal']}", ""]
    if g.get("deliverable"):
        L += [f"**You hand in:** {g['deliverable']}", ""]
    if g.get("total_time"):
        L += [f"**Time:** {g['total_time']}", ""]
    if g.get("why_it_matters"):
        L += [g["why_it_matters"], ""]
    if g.get("confirm_first"):
        L += ["## Ask before you start", ""] + [f"- {q}" for q in g["confirm_first"]] + [""]
    if g.get("requirements"):
        L += ["## What you'll need", ""] + [f"- **{r['item']}**" + (f" — {r['why']}" if r.get("why") else "") for r in g["requirements"]] + [""]
    L += ["## Steps", ""]
    for n, s in enumerate(g.get("steps", []), 1):
        mark = "x" if s["id"] in done else " "
        mins = f" (about {s['minutes']} min)" if s.get("minutes") else ""
        L += [f"### {n}. {s['title']}{mins}", "", f"- [{mark}] {s['what_to_do']}" if s.get("what_to_do") else f"- [{mark}] {s['title']}"]
        if s.get("how"):
            L += [f"- **How:** {s['how']}"]
        if s.get("done_when"):
            L += [f"- **Done when:** {s['done_when']}"]
        if s.get("watch_out"):
            L += [f"- **Watch out:** {s['watch_out']}"]
        L += [""]
    if g.get("quality_check"):
        L += ["## Check before you hand it in", ""] + [f"- [ ] {q}" for q in g["quality_check"]] + [""]
    if g.get("common_mistakes"):
        L += ["## Common mistakes", ""] + [f"- {q}" for q in g["common_mistakes"]] + [""]
    if include_materials:
        for m in g.get("materials", []):
            L += ["---", "", f"## Template: {m['name']}", ""]
            if m.get("purpose"):
                L += [f"_{m['purpose']}_", ""] if False else [m["purpose"], ""]
            body = m.get("template", "").strip()
            # keep the template's own headings below the guide's
            body = re.sub(r"^(#{1,3})\s", lambda mo: "#" * min(4, len(mo.group(1)) + 2) + " ", body, flags=re.M)
            L += [body, ""]
    return "\n".join(L).strip() + "\n"


def share_text(g: dict, done: list[str] | None = None) -> str:
    """Plain text for chat, email or a notes app."""
    done = set(done or [])
    L = [g.get("title") or "Task guide"]
    if g.get("goal"):
        L += ["", g["goal"]]
    if g.get("deliverable"):
        L += [f"You hand in: {g['deliverable']}"]
    L += ["", "STEPS"]
    for n, s in enumerate(g.get("steps", []), 1):
        L += [f"{'[x]' if s['id'] in done else '[ ]'} {n}. {s['title']}" + (f" (about {s['minutes']} min)" if s.get("minutes") else "")]
        if s.get("done_when"):
            L += [f"      Done when: {s['done_when']}"]
    if g.get("quality_check"):
        L += ["", "CHECK BEFORE YOU HAND IT IN"] + [f"[ ] {q}" for q in g["quality_check"]]
    return "\n".join(L)
